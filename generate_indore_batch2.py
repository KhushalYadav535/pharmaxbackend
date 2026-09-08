import openpyxl
import json
import re

wb = openpyxl.load_workbook(r'd:\pharmax\DOCTOR List Format (1).xlsx')
sheet = wb.active

def clean_name(raw_name):
    # Remove 'Dr.' or 'Dr ' prefix if any
    name = re.sub(r'^(Dr\.?|DR\.?)\s+', '', raw_name.strip(), flags=re.IGNORECASE)
    parts = [p.capitalize() if not p.isupper() or len(p) > 2 else p.upper() for p in name.split()]
    # Title-case each word unless it's initials like 'M K' or 'S S'
    formatted_parts = []
    for p in name.split():
        if len(p) <= 2 and p.isalpha():
            formatted_parts.append(p.upper())
        else:
            formatted_parts.append(p.capitalize())
            
    if len(formatted_parts) == 1:
        return formatted_parts[0], None, "Kumar"
    elif len(formatted_parts) == 2:
        return formatted_parts[0], None, formatted_parts[1]
    elif len(formatted_parts) == 3:
        return formatted_parts[0], formatted_parts[1], formatted_parts[2]
    else:
        return formatted_parts[0], " ".join(formatted_parts[1:-1]), formatted_parts[-1]

def clean_phone(raw_phone):
    if not raw_phone:
        return None
    s = str(raw_phone).strip()
    if s.endswith('.0'):
        s = s[:-2]
    # Remove non-digits
    digits = re.sub(r'\D', '', s)
    if len(digits) == 10:
        return digits
    elif len(digits) == 11 and digits.startswith('0'):
        return digits[1:]
    elif len(digits) == 12 and digits.startswith('91'):
        return digits[2:]
    elif len(digits) > 10:
        return digits[:10]
    return digits if digits else None

def map_specialty_qualification(desig):
    d = (desig or '').strip().upper()
    if 'PAEDIATRICIAN' in d or 'DCH' in d:
        return {
            'specialty': 'Pediatrics',
            'qualification': 'MBBS, MD (Pediatrics)' if 'PAEDIATRICIAN' in d else 'MBBS, DCH',
            'classification': 'A',
            'category': 'Pediatrician',
            'prescriptionPotential': 40000
        }
    if 'CONSULTING PHYSICIAN' in d:
        return {
            'specialty': 'General Medicine',
            'qualification': 'MBBS, MD (Medicine)',
            'classification': 'A_PLUS',
            'category': 'Consultant Physician',
            'prescriptionPotential': 50000
        }
    if 'GYNAECOLOGIST' in d:
        return {
            'specialty': 'Gynecology',
            'qualification': 'MBBS, MS (OBG)',
            'classification': 'A',
            'category': 'Consultant Gynecologist',
            'prescriptionPotential': 45000
        }
    if 'MBBS, MD' in d or d == 'MD':
        return {
            'specialty': 'General Medicine',
            'qualification': 'MBBS, MD',
            'classification': 'A_PLUS',
            'category': 'Consultant Physician',
            'prescriptionPotential': 45000
        }
    if 'MBBS' in d:
        return {
            'specialty': 'General Physician',
            'qualification': 'MBBS',
            'classification': 'B',
            'category': 'General Physician',
            'prescriptionPotential': 30000
        }
    if 'BAMS' in d:
        return {
            'specialty': 'General Practitioner',
            'qualification': 'BAMS',
            'classification': 'B',
            'category': 'General Practitioner (RMP)',
            'prescriptionPotential': 25000
        }
    if 'BHMS' in d:
        return {
            'specialty': 'General Practitioner',
            'qualification': 'BHMS',
            'classification': 'B',
            'category': 'General Practitioner (RMP)',
            'prescriptionPotential': 25000
        }
    if 'BEMS' in d:
        return {
            'specialty': 'General Practitioner',
            'qualification': 'BEMS',
            'classification': 'B',
            'category': 'General Practitioner (RMP)',
            'prescriptionPotential': 20000
        }
    return {
        'specialty': 'General Physician',
        'qualification': 'MBBS',
        'classification': 'B',
        'category': 'General Physician',
        'prescriptionPotential': 30000
    }

def get_area_and_location(area_str, addr_str):
    a_str = (area_str or '').strip().upper()
    addr_clean = (addr_str or '').strip()
    addr_lower = addr_clean.lower()
    
    if 'DEWAS' in a_str:
        area_name = 'Dewas'
        city = 'Dewas'
        district = 'Dewas'
        pincode = '455001'
        full_addr = f"{addr_clean}, Dewas, MP - {pincode}" if addr_clean else f"Dewas, Madhya Pradesh - {pincode}"
        addr1 = addr_clean if addr_clean else "Dewas, Madhya Pradesh"
        return area_name, city, district, pincode, full_addr, addr1

    if 'RAU' in a_str or 'RANGVASA' in a_str:
        area_name = 'Rau - Rangwasa'
        city = 'Indore'
        district = 'Indore'
        pincode = '453331'
        full_addr = f"{addr_clean}, Rau, Indore, MP - {pincode}" if addr_clean else f"Rau, Indore, Madhya Pradesh - {pincode}"
        addr1 = addr_clean if addr_clean else "Rau, Indore"
        return area_name, city, district, pincode, full_addr, addr1

    # Indore City addresses
    city = 'Indore'
    district = 'Indore'
    
    if any(k in addr_lower for k in ['khajrana', 'shree nagar', 'heena palace']):
        area_name = 'Khajrana'
        pincode = '452016'
    elif any(k in addr_lower for k in ['tilak nagar', 'bangali', 'bengali', 'kanadia', 'sanvid', 'saket', 'sanket', 'goyal nagar', 'patrakar', 'anand bazar', 'ananda bazar', 'pipliya', 'bhaktaver', 'bicholi', 'sanchar nagar']):
        area_name = 'Bengali Square - Tilak Nagar'
        pincode = '452016'
    elif any(k in addr_lower for k in ['phoenix', 'narayani', 'sector a', 'udhyog', 'telephone square', 'vijay']):
        area_name = 'Vijay Nagar'
        pincode = '452010'
    else:
        area_name = 'Palasia - Geeta Bhawan'
        pincode = '452001'

    full_addr = f"{addr_clean}, Indore, MP - {pincode}" if addr_clean else f"{area_name}, Indore, Madhya Pradesh - {pincode}"
    addr1 = addr_clean if addr_clean else f"{area_name}, Indore"
    return area_name, city, district, pincode, full_addr, addr1

doctors = []
for r in range(3, sheet.max_row + 1):
    sn = sheet.cell(row=r, column=1).value
    raw_name = sheet.cell(row=r, column=2).value
    desig = sheet.cell(row=r, column=3).value
    area = sheet.cell(row=r, column=4).value
    phone = sheet.cell(row=r, column=6).value
    addr = sheet.cell(row=r, column=7).value
    
    if not raw_name or not str(raw_name).strip():
        continue
        
    doc_index = len(doctors) + 1 # 1 to 100
    code_num = 100 + doc_index    # 101 to 200
    doctor_code = f"DOC-IND-{code_num:03d}"
    
    first_name, middle_name, last_name = clean_name(str(raw_name))
    spec_info = map_specialty_qualification(desig)
    cleaned_ph = clean_phone(phone)
    area_name, city, district, pincode, full_addr, addr1 = get_area_and_location(area, addr)
    
    doc_record = {
        "sn": doc_index,
        "doctorCode": doctor_code,
        "salutation": "Dr.",
        "firstName": first_name,
        "middleName": middle_name,
        "lastName": last_name,
        "specialty": spec_info['specialty'],
        "qualification": spec_info['qualification'],
        "classification": spec_info['classification'],
        "category": spec_info['category'],
        "prescriptionPotential": spec_info['prescriptionPotential'],
        "prescriber": True,
        "area": area_name,
        "address": full_addr,
        "address1": addr1,
        "city": city,
        "district": district,
        "state": "Madhya Pradesh",
        "pincode": pincode,
        "phone": cleaned_ph,
        "whatsappNumber": cleaned_ph,
        "email": None
    }
    doctors.append(doc_record)

print(f"Generated {len(doctors)} doctor records.")
print("Sample record 1 (DOC-IND-101):", json.dumps(doctors[0], indent=2))
print("Sample record 50 (DOC-IND-150):", json.dumps(doctors[49], indent=2))
print("Sample record 85 (Dewas - DOC-IND-185):", json.dumps(doctors[84], indent=2))
print("Sample record 95 (Rau - DOC-IND-195):", json.dumps(doctors[94], indent=2))
print("Sample record 100 (DOC-IND-200):", json.dumps(doctors[99], indent=2))

output_path = r'd:\pharmax\backend\indore_doctors_batch2.json'
with open(output_path, 'w', encoding='utf-8') as f:
    json.dump(doctors, f, indent=2, ensure_ascii=False)

print(f"Successfully saved to {output_path}")
