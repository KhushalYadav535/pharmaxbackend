import openpyxl
import json
import re

AREA_CONFIG = {
    'Wardha': {
        'name': 'Wardha',
        'code': 'AREA-WARD-CITY',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442001'
    },
    'Hinganghat': {
        'name': 'Hinganghat',
        'code': 'AREA-WARD-HINGANGHAT',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442301'
    },
    'Arvi': {
        'name': 'Arvi',
        'code': 'AREA-WARD-ARVI',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442201'
    },
    'Pulgaon + Rohna': {
        'name': 'Pulgaon - Rohna',
        'code': 'AREA-WARD-PULGAON',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442302'
    },
    'Deoli': {
        'name': 'Deoli',
        'code': 'AREA-WARD-DEOLI',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442101'
    },
    'Girad + Samudrapur': {
        'name': 'Samudrapur - Girad',
        'code': 'AREA-WARD-SAMUDRAPUR',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442905'
    },
    'Sindhi + Kelzar': {
        'name': 'Sindhi - Kelzar',
        'code': 'AREA-WARD-SINDHI',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442107'
    },
    'Selu + Hingni': {
        'name': 'Selu - Hingni',
        'code': 'AREA-WARD-SELU',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442104'
    },
    'Alipur + Wadner': {
        'name': 'Alipur - Wadner',
        'code': 'AREA-WARD-ALIPUR',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442304'
    },
    'Wadner': {
        'name': 'Wadner',
        'code': 'AREA-WARD-WADNER',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442307'
    },
    'Karmya': {
        'name': 'Karmya',
        'code': 'AREA-WARD-KARMYA',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442203'
    },
    'Nandore + Kora': {
        'name': 'Nandore - Kora',
        'code': 'AREA-WARD-KORA',
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442907'
    }
}

def clean_name(raw_name):
    s = str(raw_name).strip()
    is_mrs = bool(re.search(r'\bmrs\.?\b', s, re.IGNORECASE))
    is_mr = bool(re.search(r'\bmr\.?\b', s, re.IGNORECASE))
    
    # Check parenthesized format e.g. "Borkar (Rajendra)" -> "Rajendra Borkar"
    m_paren = re.match(r'^(?:Dr\.?|DR\.?|dr\.?)\s*(?:Mrs\.?|mrs\.?)?\s*([^\(\)]+)\s*\((.+)\)$', s, flags=re.IGNORECASE)
    if m_paren:
        first_part = m_paren.group(1).strip()
        paren_part = m_paren.group(2).strip()
        # e.g. "Vasundhara Deshpande (Ghorpade)" -> "Vasundhara", "Deshpande (Ghorpade)"
        if ' ' in first_part:
            fp_tokens = first_part.split()
            fn = fp_tokens[0]
            ln = f"{' '.join(fp_tokens[1:])} ({paren_part})"
            salutation = 'Dr. Mrs.' if is_mrs else 'Dr.'
            return salutation, fn, None, ln
        else:
            salutation = 'Dr. Mrs.' if is_mrs else 'Dr.'
            return salutation, paren_part, None, first_part

    # Strip prefixes
    clean = re.sub(r'^(?:Dr\.?|DR\.?|dr\.?)\s*', '', s, flags=re.IGNORECASE).strip()
    clean = re.sub(r'^(?:Mrs\.?|mrs\.?)\s*', '', clean, flags=re.IGNORECASE).strip()
    clean = re.sub(r'^(?:Mr\.?|mr\.?)\s*', '', clean, flags=re.IGNORECASE).strip()

    tokens = [t.strip() for t in clean.split() if t.strip()]
    formatted = []
    for t in tokens:
        if '.' in t:
            sub = t.split('.')
            clean_sub = '.'.join([s.capitalize() if len(s) > 1 else s.upper() for s in sub])
            formatted.append(clean_sub)
        elif len(t) <= 2 and t.isalpha():
            formatted.append(t.upper())
        else:
            formatted.append(t.capitalize())

    salutation = 'Dr. Mrs.' if is_mrs else 'Dr.'

    if len(formatted) == 0:
        return salutation, 'Doctor', None, 'Wardha'
    elif len(formatted) == 1:
        # e.g. "Dr. Mrs. Borkar" -> firstName: "Mrs.", lastName: "Borkar"
        if is_mrs:
            return 'Dr.', 'Mrs.', None, formatted[0]
        elif is_mr:
            return 'Dr.', 'Mr.', None, formatted[0]
        else:
            return 'Dr.', formatted[0], None, ''
    elif len(formatted) == 2:
        if is_mrs:
            return 'Dr. Mrs.', formatted[0], None, formatted[1]
        elif is_mr:
            return 'Dr. Mr.', formatted[0], None, formatted[1]
        else:
            return 'Dr.', formatted[0], None, formatted[1]
    elif len(formatted) == 3:
        if is_mrs:
            return 'Dr. Mrs.', formatted[0], formatted[1], formatted[2]
        return 'Dr.', formatted[0], formatted[1], formatted[2]
    else:
        if is_mrs:
            return 'Dr. Mrs.', formatted[0], ' '.join(formatted[1:-1]), formatted[-1]
        return 'Dr.', formatted[0], ' '.join(formatted[1:-1]), formatted[-1]


def map_specialty_qualification(raw_sp, exp_sp):
    raw = (raw_sp or '').strip().upper()
    exp = (exp_sp or '').strip().upper()
    combo = f"{raw} {exp}"

    if 'GYM' in combo or 'GYNAEC' in combo or 'GYNEC' in combo:
        return {
            'specialty': 'Gynecology & Obstetrics',
            'qualification': 'DGO / MS (OBG)',
            'classification': 'A',
            'category': 'Consultant Gynecologist',
            'prescriptionPotential': 45000
        }
    elif 'PIDIA' in combo or 'PAEDIATRIC' in combo or 'PEDIATRIC' in combo:
        return {
            'specialty': 'Pediatrics',
            'qualification': 'MD (Pediatrics)',
            'classification': 'A',
            'category': 'Pediatrician',
            'prescriptionPotential': 40000
        }
    elif 'ORTHO' in combo:
        return {
            'specialty': 'Orthopedics',
            'qualification': 'MBBS, MS (Ortho)',
            'classification': 'A',
            'category': 'Consultant Orthopedic',
            'prescriptionPotential': 45000
        }
    elif 'PHY' in combo or 'PHYSICIAN' in combo:
        return {
            'specialty': 'General Medicine',
            'qualification': 'MBBS, MD (Medicine)',
            'classification': 'A_PLUS',
            'category': 'Consultant Physician',
            'prescriptionPotential': 50000
        }
    elif 'SURG' in combo:
        return {
            'specialty': 'General Surgery',
            'qualification': 'MBBS, MS (General Surgery)',
            'classification': 'A',
            'category': 'Consultant Surgeon',
            'prescriptionPotential': 45000
        }
    elif 'ENT' in combo:
        return {
            'specialty': 'ENT',
            'qualification': 'MBBS, MS (ENT)',
            'classification': 'A',
            'category': 'ENT Specialist',
            'prescriptionPotential': 40000
        }
    elif 'BAMS' in combo:
        return {
            'specialty': 'General Practitioner',
            'qualification': 'BAMS',
            'classification': 'B',
            'category': 'General Practitioner (RMP)',
            'prescriptionPotential': 25000
        }
    elif 'DHMS' in combo:
        return {
            'specialty': 'General Practitioner',
            'qualification': 'DHMS',
            'classification': 'C',
            'category': 'General Practitioner (RMP)',
            'prescriptionPotential': 20000
        }
    elif 'MBBS' in combo:
        return {
            'specialty': 'General Medicine',
            'qualification': 'MBBS',
            'classification': 'B',
            'category': 'General Physician',
            'prescriptionPotential': 35000
        }
    else: # GP or None or General Practitioner
        return {
            'specialty': 'General Medicine',
            'qualification': 'MBBS / GP',
            'classification': 'B',
            'category': 'General Physician',
            'prescriptionPotential': 30000
        }

wb = openpyxl.load_workbook(r'd:\pharmax\Doctors List_Wardha_Checked.xlsx')
sheet = wb['Doctors']

doctors = []
code_index = 1

for r in range(2, sheet.max_row + 1):
    raw_area = sheet.cell(r, 1).value
    raw_sr = sheet.cell(r, 2).value
    raw_name = sheet.cell(r, 3).value
    raw_sp = sheet.cell(r, 4).value
    exp_sp = sheet.cell(r, 5).value

    if not raw_name or not str(raw_name).strip():
        continue

    area_key = str(raw_area).strip()
    area_meta = AREA_CONFIG.get(area_key, {
        'name': area_key,
        'code': f"AREA-WARD-{code_index:03d}",
        'district': 'Wardha',
        'state': 'Maharashtra',
        'pinCode': '442001'
    })

    salutation, fn, mn, ln = clean_name(raw_name)
    spec_meta = map_specialty_qualification(raw_sp, exp_sp)

    doc_code = f"DOC-WARD-{str(code_index).zfill(3)}"
    code_index += 1

    address_val = f"{area_meta['name']}, Wardha District, Maharashtra"

    doctors.append({
        'sn': code_index - 1,
        'doctorCode': doc_code,
        'salutation': salutation,
        'firstName': fn,
        'middleName': mn,
        'lastName': ln,
        'specialty': spec_meta['specialty'],
        'qualification': spec_meta['qualification'],
        'classification': spec_meta['classification'],
        'category': spec_meta['category'],
        'area': area_meta['name'],
        'areaCode': area_meta['code'],
        'city': 'Wardha' if area_key == 'Wardha' else area_meta['name'],
        'district': area_meta['district'],
        'state': area_meta['state'],
        'pincode': area_meta['pinCode'],
        'phone': None,
        'whatsappNumber': None,
        'email': None,
        'address': address_val,
        'address1': address_val,
        'prescriber': True,
        'prescriptionPotential': spec_meta['prescriptionPotential'],
        'visitFrequency': 2,
        'approvalStatus': 'APPROVED',
        'isActive': True
    })

print(f"Total parsed Wardha doctors: {len(doctors)}")

out_path = r'd:\pharmax\backend\wardha_doctors.json'
with open(out_path, 'w', encoding='utf-8') as f:
    json.dump(doctors, f, indent=2, ensure_ascii=False)

print(f"[SUCCESS] Successfully generated and saved to {out_path}")
