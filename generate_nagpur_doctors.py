import openpyxl
import json
import re

def clean_name(raw):
    s = re.sub(r'^(?:Dr\.?|DR\.?|dr\.?)\s*', '', str(raw).strip())
    parts = s.split()
    parts = [p.capitalize() for p in parts if p]
    if len(parts) == 0:
        return 'Unknown', 'Doctor'
    elif len(parts) == 1:
        return parts[0], ''
    elif len(parts) == 2:
        return parts[0], parts[1]
    else:
        return parts[0], ' '.join(parts[1:])

wb = openpyxl.load_workbook(r'd:\pharmax\DOCTOR List Format vijay.xlsx')
sheet = wb.active

doctors = []
code_index = 1

for r in range(3, sheet.max_row + 1):
    raw_sn = sheet.cell(r, 1).value
    raw_name = sheet.cell(r, 2).value
    raw_desig = sheet.cell(r, 3).value
    raw_area = sheet.cell(r, 4).value
    raw_email = sheet.cell(r, 5).value
    raw_mobile = sheet.cell(r, 6).value
    raw_addr = sheet.cell(r, 7).value

    if not raw_name or not str(raw_name).strip():
        continue

    fn, ln = clean_name(raw_name)
    desig_str = str(raw_desig).strip() if raw_desig else ''
    desig_upper = desig_str.upper()

    # Specialty & Category mapping
    if 'DGO' in desig_upper or 'GYN' in desig_upper:
        specialty = 'Gynecology & Obstetrics'
        category = 'Consultant Gynecologist'
        qualification = 'DGO'
    elif 'PAED' in desig_upper or 'PED' in desig_upper:
        specialty = 'Pediatrics'
        category = 'Pediatrician'
        qualification = 'MD (Pediatrics)'
    elif 'GP' in desig_upper or 'MBBS' in desig_upper:
        specialty = 'General Medicine'
        category = 'General Physician'
        qualification = 'MBBS'
    else:
        specialty = 'General Medicine'
        category = 'General Physician'
        qualification = desig_str or 'MBBS'

    doc_code = f"DOC-NAGP-{str(code_index).zfill(3)}"
    code_index += 1

    address_val = str(raw_addr).strip() if raw_addr and str(raw_addr).strip() != 'None' else 'Nagpur, Maharashtra'
    mobile_val = str(raw_mobile).strip() if raw_mobile and str(raw_mobile).strip() != 'None' else None
    email_val = str(raw_email).strip() if raw_email and str(raw_email).strip() != 'None' else None

    doctors.append({
        'sn': code_index - 1,
        'doctorCode': doc_code,
        'salutation': 'Dr.',
        'firstName': fn,
        'middleName': None,
        'lastName': ln,
        'specialty': specialty,
        'qualification': qualification,
        'classification': 'A',
        'category': category,
        'area': 'Nagpur Central Area',
        'city': 'Nagpur',
        'district': 'Nagpur',
        'state': 'Maharashtra',
        'pincode': '440001',
        'phone': mobile_val,
        'whatsappNumber': mobile_val,
        'email': email_val,
        'address': address_val,
        'address1': address_val,
        'prescriber': True,
        'prescriptionPotential': 40000,
        'visitFrequency': 2,
        'approvalStatus': 'APPROVED',
        'isActive': True
    })

print(f"Total parsed doctors: {len(doctors)}")
with open(r'd:\pharmax\backend\nagpur_doctors.json', 'w', encoding='utf-8') as f:
    json.dump(doctors, f, indent=2, ensure_ascii=False)

print("Saved to d:\\pharmax\\backend\\nagpur_doctors.json successfully.")
